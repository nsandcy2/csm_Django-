import os, tempfile, pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from django.shortcuts import render, redirect
from django.http import JsonResponse, FileResponse
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError

from .models import User, Project

# Global temp storage for uploaded Excel files
UPLOAD_FOLDER = tempfile.gettempdir()
excel_data = {}

from django.urls import reverse



def home(request):
    if request.user.is_authenticated:
        return redirect("dashboard")   # Django uses URL names, not function names`
    return redirect("login")

# ------------------ Auth ------------------
def register(request):
    if request.method == "POST":
        username = request.POST["username"].strip()
        password = request.POST["password"]
        role = request.POST.get("role", "user")

        if not username or not password:
            messages.error(request, "Username and password are required")
            return redirect("register")

        try:
            User.objects.create_user(username=username, password=password, role=role)
            messages.success(request, "Registration successful! Please log in.")
            return redirect("login")
        except IntegrityError:
            messages.warning(request, "Username already exists")
            return redirect("register")

    return render(request, "core/register.html")


def login_view(request):
    if request.method == "POST":
        username = request.POST["username"]
        password = request.POST["password"]

        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            messages.success(request, "Login successful!")
            if user.role == "admin":
                return redirect("dashboard")
            return redirect("user_dashboard")
        else:
            messages.error(request, "Invalid username or password")

    return render(request, "core/login.html")


@login_required
def logout_view(request):
    logout(request)
    messages.info(request, "Logged out successfully!")
    return redirect("login")


# ------------------ Dashboards ------------------
@login_required
def dashboard(request):
    if request.user.role != "admin":
        return redirect("user_dashboard")
    projects = Project.objects.all()
    return render(request, "core/dashboard.html", {"username": request.user.username, "projects": projects})


@login_required
def user_dashboard(request):
    projects = Project.objects.filter(owner=request.user)
    return render(request, "core/user_dashboard.html", {"username": request.user.username, "projects": projects})


# ------------------ Projects ------------------
@login_required
def add_project(request):
    if request.method == "POST":
        raw = request.POST.get("project_name", "").strip()
        name = raw or f"Project {Project.objects.count() + 1}"

        if Project.objects.filter(name=name).exists():
            messages.error(request, "Project name already exists")
            return redirect("dashboard")

        Project.objects.create(name=name, created_by=request.user)
        messages.success(request, f'Project “{name}” added!')
        return redirect("dashboard")

    return render(request, "core/add_project.html")


@login_required
def assign_project(request):
    if request.method == "POST" and request.user.role == "admin":
        project_id = request.POST.get("project_id")
        user_id = request.POST.get("user_id")
        try:
            project = Project.objects.get(id=project_id)
            user = User.objects.get(id=user_id)
            project.owner = user
            project.save()
            messages.success(request, "✅ Project assigned successfully!")
        except Exception as e:
            messages.error(request, f"❌ Error: {e}")

    return render(
        request,
        "core/assign_project.html",
        {"projects": Project.objects.all(), "users": User.objects.all()},
    )


@login_required
def work_products(request, project_id):
    return render(request, "core/work_products.html", {"project_id": project_id})



@login_required
def index(request):
    return render(request, "core/login.html")


# ------------------ Excel Handling ------------------
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods


@csrf_exempt
@require_http_methods(["POST"])
def upload(request):
    file = request.FILES["file"]
    if file.name.endswith((".xlsx", ".xls")):
        filepath = os.path.join(UPLOAD_FOLDER, file.name)
        with open(filepath, "wb+") as dest:
            for chunk in file.chunks():
                dest.write(chunk)
        excel_data[file.name] = {"path": filepath}
        xl = pd.ExcelFile(filepath)
        return JsonResponse({"message": "Uploaded", "filename": file.name, "sheets": xl.sheet_names})
    return JsonResponse({"error": "Invalid file format"}, status=400)


@require_http_methods(["GET"])
def edit(request):
    filename = request.GET.get("filename")
    sheet = request.GET.get("sheet")
    file_info = excel_data.get(filename)

    if not file_info or not os.path.exists(file_info["path"]):
        return JsonResponse({"error": "File not found"}, status=404)

    df = pd.read_excel(file_info["path"], sheet_name=sheet, dtype=str).fillna("")
    wb = load_workbook(file_info["path"], data_only=True)
    ws = wb[sheet]

    dropdowns = {}
    if ws.data_validations:
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list" and dv.formula1:
                options = []
                if dv.formula1.startswith("="):
                    try:
                        ref = dv.formula1.strip("=").replace("$", "")
                        if "!" in ref:
                            sheetname, ref = ref.split("!")
                            target_ws = wb[sheetname]
                        else:
                            target_ws = ws
                        min_col, min_row, max_col, max_row = range_boundaries(ref)
                        for row in target_ws.iter_rows(min_row=min_row, max_row=max_row,
                                                       min_col=min_col, max_col=max_col):
                            for cell in row:
                                if cell.value:
                                    options.append(str(cell.value))
                    except Exception:
                        options = []
                else:
                    options = dv.formula1.strip('"').split(",")

                for cell_range in dv.sqref.ranges:
                    min_col, min_row, max_col, max_row = range_boundaries(str(cell_range))
                    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                                            min_col=min_col, max_col=max_col):
                        for cell in row:
                            dropdowns[cell.coordinate] = options

    return JsonResponse({
        "columns": df.columns.tolist(),
        "data": df.to_dict(orient="records"),
        "dropdowns": dropdowns
    })


@csrf_exempt
@require_http_methods(["POST"])
def save(request):
    import json
    data = json.loads(request.body)
    filename = data.get("filename")
    sheet = data.get("sheet")
    edited_data = data.get("data")

    if filename not in excel_data:
        return JsonResponse({"error": "File not found"}, status=404)

    filepath = excel_data[filename]["path"]
    df = pd.DataFrame(edited_data)

    wb = load_workbook(filepath)
    ws = wb[sheet]

    # Clear existing rows (except header)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.value = None

    for r, row_data in enumerate(df.values, start=2):
        for c, value in enumerate(row_data, start=1):
            ws.cell(row=r, column=c, value=value)

    wb.save(filepath)
    return JsonResponse({"message": "Saved successfully"})


@require_http_methods(["GET"])
def download(request):
    filename = request.GET.get("filename")
    custom_name = request.GET.get("custom_name", "Edited_File.xlsx")
    if filename in excel_data:
        filepath = excel_data[filename]["path"]
        return FileResponse(open(filepath, "rb"), as_attachment=True, filename=custom_name)
    return JsonResponse({"error": "File not found"}, status=404)
